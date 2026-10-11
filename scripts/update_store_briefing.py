"""update_store_briefing.py -- daily refresh of store-briefing.html's BRIEFING_DATA, including the
per-store "What to do today" call-outs. Run by daily_refresh.py after the footfall fetch.

Usage (from the repo root):
  python scripts/update_store_briefing.py --html store-briefing.html --sales <mtd csv>
      [--daywise-targets "Daywise October target.xlsx" --store-targets "OND overall target.xlsx"]
      [--compliance "Compliance Score Sheet - ....xlsx"]

Sources:
  - Sales: the month-to-date Shopify export (same file as the dashboards).
  - Targets: loaded ONCE per month with --daywise-targets (store + stylist revenue target per day)
    and --store-targets (monthly sheet: New Rev %, New/Repeat AOV, New/Rpt Cust, NC Conversion,
    New FF), then kept in data/briefing_targets.json so the daily run needs no xlsx.
  - Footfall: data/footfall_history_daily.csv (footfall.py fetch).
  - NPS and last month's ASP/UPT: index.html. Audit score: --compliance, else carried forward
    from the page's previous BRIEFING_DATA (it keeps its own period label).

Target basis (same as the Store Pulse page, so the two can be cross-checked):
  month revenue target = the day-wise file's Store Total row; New/Repeat revenue targets = that
  total split by the monthly sheet's New Rev %; bill targets = New Cust / Rpt Cust; conversion
  target = NC Conversion, on new-customer footfall.

Call-outs: every store gets today's target and required run-rate, then up to three "focus" items
picked from four store-team levers (new-customer conversion, New AOV, repeat customers, Repeat
AOV) -- only the ones behind target, ranked by the revenue they have cost month to date -- and a
stylist line naming who is behind / ahead of their own day-wise target. Nothing is shown that the
store cannot act on today (see BSC_Dashboard_Runbook.md, briefing conventions).
"""
import argparse
import calendar
import csv
import datetime
import difflib
import json
import math
import os

import openpyxl

from bsc_common import REGION_MAP, load_sales_csv, replace_const, extract_const, syntax_check_html_js
from generate_briefing_data import build_stylist_breakdown, load_compliance, load_index_context

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TARGETS_PATH = os.path.join(ROOT, 'data', 'briefing_targets.json')
FOOTFALL_PATH = os.path.join(ROOT, 'data', 'footfall_history_daily.csv')
MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
STORE_TARGET_ALIASES = {'MOI': 'Mall of India, Noida'}
MAX_FOCUS = 3
MIN_MISSED_BILLS = 3         # conversion is a call-out only once the gap is worth this many bills
MIN_BILLS_FOR_AOV = 5        # fewer bills than this and an AOV gap is noise, not a call-out
STYLIST_BEHIND = 0.8         # below 80% of own target-to-date = named as behind
UNTAGGED_SHARE = 0.08        # share of store revenue with no stylist on the bill worth calling out


def inr(n):
    n = int(round(abs(n)))
    s = str(n)
    head, tail = s[:-3], s[-3:]
    parts = []
    while len(head) > 2:
        parts.insert(0, head[-2:])
        head = head[:-2]
    if head:
        parts.insert(0, head)
    return 'Rs' + ','.join(parts + [tail])


def pct(x):
    return '%.0f%%' % (x * 100)


def load_daywise_targets(path):
    """(month 'YYYY-MM', {store: {'store': {date: target}, 'stylists': {name: {date: target}}}})."""
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    date_cols = [(i, c) for i, c in enumerate(rows[0]) if isinstance(c, (datetime.datetime, datetime.date))]
    if not date_cols:
        raise ValueError(f'No date columns in the header row of {path}')
    out = {}
    for r in rows[1:]:
        store, staff = r[1], r[2]
        if store is None or not staff:
            continue
        store = str(store).strip()
        series = {d.strftime('%Y-%m-%d'): round(float(r[i] or 0), 2) for i, d in date_cols}
        s = out.setdefault(store, {'store': None, 'stylists': {}})
        if str(staff).strip().lower() == 'store total':
            s['store'] = series
        else:
            s['stylists'][str(staff).strip()] = series
    unknown = [s for s in out if s not in REGION_MAP]
    if unknown:
        raise ValueError(f'Stores in {path} not in REGION_MAP: {unknown}')
    return date_cols[0][1].strftime('%Y-%m'), {s: v for s, v in out.items() if v['store']}


def load_store_targets(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = list(wb[sheet].iter_rows(values_only=True))
    hdr = [str(c).strip().lower() if c is not None else '' for c in rows[0]]
    keys = {'newRevShare': 'new rev %', 'newAov': 'new aov', 'repAov': 'repeat aov', 'newCust': 'new cust',
            'repCust': 'rpt cust', 'convTarget': 'nc conversion', 'newFF': 'new ff'}
    col = {k: hdr.index(v) for k, v in keys.items()}
    name_i = hdr.index('store name')
    out = {}
    for r in rows[1:]:
        if r[name_i] is None:
            continue
        name = STORE_TARGET_ALIASES.get(str(r[name_i]).strip(), str(r[name_i]).strip())
        if name not in REGION_MAP:
            print(f"[warn] '{name}' in {sheet} store targets is not a known store -- skipped")
            continue
        out[name] = {k: (float(r[i]) if isinstance(r[i], (int, float)) else None) for k, i in col.items()}
    return out


def load_footfall(month):
    """{store: {date: (new, repeat)}} for the month; a blank cell stays None (not recorded)."""
    out = {}
    with open(FOOTFALL_PATH, encoding='utf-8') as f:
        for r in csv.DictReader(f):
            if r['date'][:7] != month:
                continue
            num = lambda v: float(v) if v not in ('', None) else None
            out.setdefault(r['store'], {})[r['date']] = (num(r['new']), num(r['repeat']))
    return out


def match_stylist(name, by_norm, taken):
    """The sales-file stylist a target-file name refers to: exact (case-insensitive), else a unique
    close spelling or a shorter form of the same name ('Pari' / 'Pari Chauhan') at the same store."""
    n = name.lower().strip()
    if n in by_norm:
        return n
    free = [k for k in by_norm if k not in taken]
    close = [k for k in free if difflib.SequenceMatcher(None, n, k).ratio() >= 0.85
             or set(n.split()) < set(k.split()) or set(k.split()) < set(n.split())]
    return close[0] if len(close) == 1 else None


def seg_stats(sub):
    rev, bills, units = float(sub['Revenue'].sum()), int(sub['Order name'].nunique()), float(sub['Qty'].sum())
    return {'rev': rev, 'bills': bills, 'units': units, 'aov': rev / bills if bills else None,
            'asp': rev / units if units else None, 'upt': units / bills if bills else None}


def segment_block(a, rev_t, bills_t, aov_t):
    """Same field names the page already reads (see generate_briefing_data.segment_block)."""
    rev_rem, bills_rem = rev_t - a['rev'], bills_t - a['bills']
    b = {
        'targetRev': round(rev_t, 2), 'targetOrders': bills_t, 'achievedRev': round(a['rev'], 2),
        'achievedBills': a['bills'], 'achievedUnits': round(a['units'], 1),
        'currentAOV': round(a['aov'], 2) if a['aov'] else None,
        'currentASP': round(a['asp'], 2) if a['asp'] else None,
        'currentUPT': round(a['upt'], 3) if a['upt'] else None,
        'hoTargetAOV': aov_t, 'remainingRev': round(rev_rem, 2), 'remainingOrders': bills_rem,
        'noTarget': rev_t == 0 and a['rev'] == 0,
    }
    if rev_rem <= 0 and not b['noTarget']:
        b['status'], b['surplus'] = 'achieved', round(-rev_rem, 2)
    elif bills_rem > 0:
        b['status'], b['requiredAOV'] = 'on-track', round(rev_rem / bills_rem, 2)
    else:
        b['status'] = 'orders-hit-revenue-short'
    return b


def aov_action(label, a, target, prev_asp, prev_upt):
    """Call-out for a segment whose AOV is under target: says which of ASP / UPT is the drag and
    what closing it means per bill."""
    gap = target - a['aov']
    head = f"{label} AOV is {inr(a['aov'])} against a {inr(target)} target, {inr(gap)} short on every bill."
    upt_need = target / a['asp']
    extra = upt_need - a['upt']
    if extra >= 1:
        add = f"about {extra:.0f} more piece{'s' if round(extra) > 1 else ''} on every bill"
    else:
        add = f"one more piece on about 1 in every {max(2, round(1 / extra))} bills"
    upt_line = (f"Bills average {a['upt']:.2f} pieces; at current prices they need {upt_need:.2f} "
                f"({add}). Offer the add-on before billing: a second shirt, trousers or an accessory.")
    asp_line = (f"Pieces are selling at {inr(a['asp'])} each" +
                (f", down from {inr(prev_asp)} last month" if prev_asp and prev_asp > a['asp'] else '') +
                ". Show the higher-value fabrics and custom options first, entry-price pieces second.")
    if prev_asp and prev_upt:
        upt_is_drag = (a['upt'] - prev_upt) / prev_upt <= (a['asp'] - prev_asp) / prev_asp
        if upt_is_drag and prev_upt > a['upt']:
            upt_line = upt_line.replace('pieces;', f"pieces (was {prev_upt:.2f} last month);", 1)
        detail = upt_line if upt_is_drag else asp_line
    else:
        detail = upt_line
    return {'key': label.lower() + 'Aov', 'title': f'Lift {label} AOV', 'detail': head + ' ' + detail,
            'impact': round(gap * a['bills'])}


def build_store(store, sub, tg, ff, ctx, last_day, d_el, d_in, prev_audit):
    iso = last_day.isoformat()
    today = (last_day + datetime.timedelta(days=1)).isoformat()
    d_left = d_in - d_el
    day_t = tg['store']
    month_t = sum(day_t.values())
    todate_t = sum(v for d, v in day_t.items() if d <= iso)
    new, rep, allseg = seg_stats(sub[sub['Segment'] == 'new']), seg_stats(sub[sub['Segment'] == 'returning']), seg_stats(sub)
    share = tg.get('newRevShare')
    has_seg = share is not None
    new_rev_t = month_t * share if has_seg else 0.0
    new_bills_t = round(tg['newCust']) if has_seg and tg.get('newCust') else 0
    rep_bills_t = round(tg['repCust']) if has_seg and tg.get('repCust') else 0
    new_block = segment_block(new, new_rev_t, new_bills_t, tg.get('newAov'))
    rep_block = segment_block(rep, month_t - new_rev_t if has_seg else 0.0, rep_bills_t, tg.get('repAov'))

    prev_upt, prev_aov = ctx['prevMonthStoreUPT'].get(store), ctx['prevMonthStoreAOV'].get(store)
    prev = None
    if prev_upt and prev_aov:
        prev = {'newUpt': prev_upt.get('newUpt'), 'repUpt': prev_upt.get('repUpt'),
                'newAsp': round(prev_aov['newAov'] / prev_upt['newUpt'], 2) if prev_upt.get('newUpt') else None,
                'repAsp': round(prev_aov['repAov'] / prev_upt['repUpt'], 2) if prev_upt.get('repUpt') else None,
                'newAov': prev_aov.get('newAov'), 'repAov': prev_aov.get('repAov')}

    # ---- new-customer conversion, only over days whose footfall was actually recorded
    new_by_day = sub[sub['Segment'] == 'new'].groupby('Day_str')['Order name'].nunique().to_dict()
    ff_days = {d: v for d, v in ff.items() if d <= iso and v[0] is not None and (v[0] or 0) + (v[1] or 0) > 0}
    ff_new = sum(v[0] for v in ff_days.values())
    conv_bills = sum(new_by_day.get(d, 0) for d in ff_days)
    conv = conv_bills / ff_new if ff_new else None
    conv_t = tg.get('convTarget')
    if conv is not None:
        new_block.update({'ffAchieved': round(ff_new, 1), 'ffDays': len(ff_days)})
        if conv > 1.0:
            new_block['ffUnreliable'] = True
            conv = None
        else:
            new_block['currentConversion'] = round(conv, 4)
    if conv_t:
        new_block['targetConversion'] = round(conv_t, 4)

    # ---- pace line (always shown) + ranked focus items
    ach = allseg['rev']
    gap = todate_t - ach
    need_day = max(month_t - ach, 0) / d_left if d_left else None
    today_t = day_t.get(today)
    pace = {'aheadBy': round(-gap), 'todayTarget': today_t, 'needPerDay': round(need_day) if need_day is not None else None,
            'runRate': round(ach / d_el)}
    y_rev = float(sub[sub['Day_str'] == iso]['Revenue'].sum())
    yesterday = {'date': iso, 'rev': round(y_rev, 2), 'target': day_t.get(iso),
                 'bills': int(sub[sub['Day_str'] == iso]['Order name'].nunique())}

    focus = []
    missed = (conv_t - conv) * ff_new if conv is not None and conv_t else 0
    if missed >= MIN_MISSED_BILLS:
        weekend = datetime.date.fromisoformat(today).weekday() >= 5
        same = [v[0] for d, v in ff_days.items() if (datetime.date.fromisoformat(d).weekday() >= 5) == weekend]
        exp = sum(same) / len(same) if same else ff_new / len(ff_days)
        detail = (f"{pct(conv)} of new walk-ins have bought ({conv_bills} of {ff_new:.0f}) against a {pct(conv_t)} target, "
                  f"about {missed:.0f} bills missed so far.")
        if exp >= 1:
            goal, now = math.ceil(exp * conv_t), round(exp * conv)
            detail += (f" Expect around {exp:.0f} new walk-ins today: bill at least {goal} of them"
                       + (f" (your current rate would give {now})" if now < goal else '')
                       + ". Greet and start a fitting with every new walk-in.")
        focus.append({'key': 'conversion', 'title': 'Convert more new walk-ins', 'detail': detail,
                      'impact': round(missed * (new['aov'] or 0))})
    if new['bills'] >= MIN_BILLS_FOR_AOV and tg.get('newAov') and new['aov'] < tg['newAov']:
        focus.append(aov_action('New', new, tg['newAov'], prev and prev['newAsp'], prev and prev['newUpt']))
    if rep_bills_t and rep['bills'] / d_el * d_in < rep_bills_t:
        need = rep_bills_t - rep['bills']
        detail = (f"{rep['bills']} repeat customers so far against a month target of {rep_bills_t}; on this pace the month ends at "
                  f"{rep['bills'] / d_el * d_in:.0f}.")
        if d_left:
            detail += (f" You need {need} more in {d_left} days, about {math.ceil(need / d_left)} a day (so far {rep['bills'] / d_el:.1f} a day)."
                       " Call past customers before the store gets busy: start with those due a reorder and anyone with an order ready for pickup.")
        focus.append({'key': 'repeat', 'title': 'Bring back repeat customers', 'detail': detail,
                      'impact': round((rep_bills_t * d_el / d_in - rep['bills']) * (rep['aov'] or tg.get('repAov') or 0))})
    if rep['bills'] >= MIN_BILLS_FOR_AOV and tg.get('repAov') and rep['aov'] < tg['repAov']:
        focus.append(aov_action('Repeat', rep, tg['repAov'], prev and prev['repAsp'], prev and prev['repUpt']))
    focus.sort(key=lambda f: -f['impact'])

    # ---- stylists against their own day-wise targets
    stylists = build_stylist_breakdown(sub)
    by_norm = {s['name'].lower(): s for s in stylists}
    taken, no_sales = set(), []
    names = sorted(tg.get('stylists', {}), key=lambda n: n.lower() not in by_norm)   # exact matches claim first
    for name in names:
        series = tg['stylists'][name]
        mtd_t = sum(v for d, v in series.items() if d <= iso)
        key = match_stylist(name, by_norm, taken)
        if key is None:
            # A target but no bills under that name: on leave, moved store, or not being tagged --
            # not something to print as "0% of target".
            if mtd_t > 0:
                no_sales.append(name)
            continue
        taken.add(key)
        by_norm[key].update({'targetToDate': round(mtd_t, 2), 'monthTarget': round(sum(series.values()), 2),
                             'todayTarget': series.get(today)})
    rated = [s for s in stylists if s.get('targetToDate')]
    behind = sorted((s for s in rated if s['totalRev'] / s['targetToDate'] < STYLIST_BEHIND), key=lambda s: s['totalRev'] / s['targetToDate'])
    ahead = sorted((s for s in rated if s['totalRev'] >= s['targetToDate']), key=lambda s: -s['totalRev'] / s['targetToDate'])
    stylist_notes = []
    if behind:
        stylist_notes.append({'kind': 'behind', 'text': 'Behind their own target to date: ' + '; '.join(
            f"{s['name']} at {pct(s['totalRev'] / s['targetToDate'])} ({inr(s['targetToDate'] - s['totalRev'])} short"
            + (f", today's target {inr(s['todayTarget'])}" if s.get('todayTarget') else '') + ')' for s in behind[:3])
            + '. Agree a number for today with each of them at the morning huddle.'})
    if ahead:
        stylist_notes.append({'kind': 'ahead', 'text': 'Ahead of target: ' + ', '.join(
            f"{s['name']} ({pct(s['totalRev'] / s['targetToDate'])})" for s in ahead[:2]) + '. Call it out in the huddle.'})
    if no_sales:
        stylist_notes.append({'kind': 'untagged', 'text': 'Has a target but no bills in their name this month: ' + ', '.join(no_sales)
                              + '. If they are on the floor, check they are being tagged on their bills.'})
    products = sub[sub['Line type'] != 'shipping'] if 'Line type' in sub.columns else sub
    untagged = float(products[products['Stylist'] == '']['Revenue'].sum())
    if ach and untagged / ach >= UNTAGGED_SHARE:
        stylist_notes.append({'kind': 'untagged', 'text': f"{inr(untagged)} of sales ({pct(untagged / ach)}) has no stylist on the bill. "
                              "Tag the assisting stylist on every bill so it counts towards their target."})

    return {
        'region': REGION_MAP[store], 'monthTarget': round(month_t, 2), 'monthTargetToDate': round(todate_t, 2),
        'achievedTotal': round(ach, 2), 'new': new_block, 'rep': rep_block, 'prevMonth': prev,
        'nps': ctx['npsByStore'].get(store), 'audit': prev_audit, 'stylists': stylists,
        'today': {'pace': pace, 'yesterday': yesterday, 'focus': focus[:MAX_FOCUS], 'stylistNotes': stylist_notes},
    }


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--html', required=True)
    ap.add_argument('--sales', required=True)
    ap.add_argument('--daywise-targets')
    ap.add_argument('--store-targets')
    ap.add_argument('--store-targets-sheet')
    ap.add_argument('--compliance')
    ap.add_argument('--index')
    a = ap.parse_args()

    all_targets = {}
    if os.path.exists(TARGETS_PATH):
        with open(TARGETS_PATH, encoding='utf-8') as f:
            all_targets = json.load(f)
    if a.daywise_targets:
        month, daywise = load_daywise_targets(a.daywise_targets)
        t = all_targets.setdefault(month, {})
        for store, v in daywise.items():
            t.setdefault(store, {}).update(v)
        if a.store_targets:
            sheet = a.store_targets_sheet or MONTHS[int(month[5:]) - 1]
            for store, v in load_store_targets(a.store_targets, sheet).items():
                if store in t:
                    t[store].update(v)
        with open(TARGETS_PATH, 'w', encoding='utf-8') as f:
            json.dump(all_targets, f, separators=(',', ':'), sort_keys=True)
        print(f'  targets loaded for {month}: {len(t)} stores')
    elif a.store_targets:
        raise SystemExit('--store-targets needs --daywise-targets (it says which month the targets are for)')

    df = load_sales_csv(a.sales)
    df = df[df['POS location name'].isin(REGION_MAP.keys())]
    days = sorted(df['Day_str'].unique())
    last_day = datetime.date.fromisoformat(days[-1])
    month = days[-1][:7]
    d_in = calendar.monthrange(last_day.year, last_day.month)[1]
    d_el = last_day.day
    targets = all_targets.get(month)
    if not targets:
        raise SystemExit(f'No targets for {month} in data/briefing_targets.json -- load them once with '
                         '--daywise-targets and --store-targets')

    with open(a.html, encoding='utf-8') as f:
        content = f.read()
    try:
        prev_stores = extract_const(content, 'BRIEFING_DATA').get('stores', {})
    except Exception:
        prev_stores = {}
    compliance = load_compliance(a.compliance) if a.compliance else {}
    index_path = a.index or os.path.join(os.path.dirname(os.path.abspath(a.html)), 'index.html')
    ctx = load_index_context(index_path)
    footfall = load_footfall(month)

    stores = {}
    for store in sorted(targets):
        sub = df[df['POS location name'] == store]
        if sub.empty:
            print(f'  [skip] {store}: has targets but no sales this month -- no briefing')
            continue
        audit = compliance.get(store) or (prev_stores.get(store) or {}).get('audit')
        stores[store] = build_store(store, sub, targets[store], footfall.get(store, {}), ctx, last_day, d_el, d_in, audit)

    payload = {
        'generatedAt': datetime.datetime.now().isoformat(timespec='seconds'),
        'lastSalesDay': last_day.isoformat(),
        'briefingFor': (last_day + datetime.timedelta(days=1)).isoformat(),
        'daysElapsed': d_el, 'daysRemaining': d_in - d_el, 'daysInMonth': d_in, 'stores': stores,
    }
    content = replace_const(content, 'BRIEFING_DATA', json.dumps(payload))
    with open(a.html, 'w', encoding='utf-8') as f:
        f.write(content)
    syntax_check_html_js(a.html)
    no_focus = [s for s, v in stores.items() if not v['today']['focus']]
    print(f"OK. {a.html}: {len(stores)} stores through {last_day} ({d_in - d_el} days left); "
          f"{len(no_focus)} with no focus item; briefing for {payload['briefingFor']}.")


if __name__ == '__main__':
    main()
