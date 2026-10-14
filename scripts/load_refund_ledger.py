#!/usr/bin/env python3
"""
load_refund_ledger.py -- Replace one month of SEED_REFUNDS in index.html from the finance
refund ledger ("Refund_Data_<range>.xlsx"). If the file also has a CX-format credit-note
sheet, that month of SEED_CN is replaced too; otherwise SEED_CN is left untouched.

The ledger is one row per refund payout, not per line item: Sr No, Order ID, Location
(region-coded, e.g. "SOU-TEL-HYD_Jubliee Hills"), Refund Amount, Taxable, GST, rate, date,
then a second block with Order ID, Location, Amount, Payment gateway, Date, Reference,
Reason. Row 1 holds column totals, row 2 the header. It has no category, so each refund is
split across categories by:
  1. the CX refund tracker's line items for that order (--cx), if the order is there;
  2. otherwise the order's sales lines (--sales-history), weighted by gross + taxes and
     mapped from Shopify product type with PRODUCT_TYPE_TO_REFUND_CAT.

Usage:
  python scripts/load_refund_ledger.py --html index.html --ledger Refund_Data_1-11_Oct_.xlsx \
      --cx "Refunds - Oct (2).xlsx" --sales-history "Sales ... 2026-01-01 - 2026-10-13.csv" \
      --month Oct-2026
"""
import argparse
import json
import re
import sys
from collections import Counter, defaultdict

import openpyxl
import pandas as pd

from bsc_common import (extract_const, find_refund_and_cn_sheets, month_key_from_date, process_refund_sheet,
                        replace_const, syntax_check_html_js)

# Ledger location code -> store name used by the CX refund/CN files (and SEED_REFUNDS.byStore).
LOCATION_MAP = {
    'Online': 'Online Store',
    'NOR-NCR-DEL_DLF Midtown': 'DLF Midtown - Moti Nagar',
    'NOR-NCR-DEL_Khan Market': 'Khan Market',
    'NOR-NCR-DEL_South Ex': 'South Ex.',
    'NOR-NCR-DEL_Vasant Kunj': 'Ambience Vasant Kunj',
    'NOR-NCR-Gurugram': 'Gurugram',
    'NOR-NCR-Saket': 'Select City',
    'NOR-NCR-UP-Phoenix Palassio': 'Phoenix Palassio',
    'NOR-NCR-UP_Mall Of India': 'Mall of India, Noida',
    'SOU-KAR-BLR_Indiranagar': 'Indiranagar',
    'SOU-TEL-HYD_Jubliee Hills': 'Jubilee Hills',
    'SOU-TEL-HYD_LakeShore Mall': 'LakeShore Mall',
    'WST-GJ-AHM_Sindhu Bhavan': 'Sindhu Bhavan Marg',
    'WST-MH-MUM_Andheri': 'Andheri',
    'WST-MH-MUM_Goregoan': 'Oberoi Mall Store',
    'WST-MH-MUM_Kala Ghoda': 'Kalaghoda, Fort',
    'WST-MH-MUM_Kemps Corner': 'Kemps Corner',
    'WST-MH-MUM_Kurla': 'Phoenix Marketcity Kurla',
    'WST-MH-MUM_Malad Inorbit mall': 'Inorbit Mall Malad West',
    'WST-MH-MUM_Pali Hill': 'Pali Hill, Bandra',
    'WST-MH-MUM_Sky City': 'Oberoi Sky City',
    'WST-MH-PUN_Koregoan Park': 'Koregaon Park',
    'WST-MH-PUN_Viman Nagar': 'Viman Nagar',
    'WST-MH-THA_Viviana Mall': 'Viviana Mall',
}

# Shopify product type -> CX refund category, as (offline, online). Derived 2026-10-14 by
# matching CX tracker line items to their sales lines (Aug-Oct): store shirts are mostly MTM,
# online shirts mostly RTW.
PRODUCT_TYPE_TO_REFUND_CAT = {
    'Shirt': ('MTM Shirt', 'RTW Shirts'),
    'Kurta': ('RTW Shirts', 'RTW Shirts'),
    'Dress Pants': ('Dress Pants', 'RTW Dress Pants'),
    'Mid Rise Dress Pants': ('Dress Pants', 'Dress Pants'),
    'Neapolitan Dress Pants': ('Neapolitan Dress Pants', 'Neapolitan Dress Pants'),
    'Relaxed Neapolitan Pants': ('Relaxed Neapolitan Pants', 'Relaxed Neapolitan Pants'),
    'Wide Leg Pants': ('Dress Pants', 'Lounge Pants'),
    'Chinos': ('Chinos', 'Chinos'),
    'Smartpants': ('Smart Pants', 'Smart Pants'),
    'Relaxed Linen Pants': ('Linen Pants', 'Linen Pants'),
    'Lounge Pants': ('Lounge Pants', 'Lounge Pants'),
    'Mid Rise Lounge Pants': ('Lounge Pants', 'Lounge Pants'),
    'Joggers': ('Jogger', 'Jogger'),
    'Mid Rise Stretch Jeans': ('Jeans', 'Jeans'),
    'Stretch Denim': ('Jeans', 'Jeans'),
    'T-Shirts': ('T-Shirts', 'T-Shirts'),
    'Relaxed T-Shirts': ('T-Shirts', 'T-Shirts'),
    'Oversized T-Shirts': ('T-Shirts', 'T-Shirts'),
    'Long Sleeve Polo T-Shirts': ('T-Shirts', 'T-Shirts'),
    'Shackets': ('Shacket', 'Shacket'),
    'Casual Blazer': ('Blazer', 'Blazer'),
    'Formal Blazer': ('Blazer', 'Blazer'),
    'Jackets': ('Blazer', 'Blazer'),
    'Soft Tailored Jackets': ('Blazer', 'Blazer'),
    'Bundi': ('Bundi', 'Bundi'),
}


REGION_CODE = re.compile(r'^(NOR|SOU|WST|EST)-')


def norm_reason(r):
    if not r or not str(r).strip():
        return 'Unspecified'
    # "FITS / FOOS" style double reasons count under the first one.
    return str(r).split('/')[0].strip().upper()


def load_ledger(path):
    """Rows of the ledger's refund sheet. Two layouts seen: Oct 1-11 had a totals row, then a
    header with Sr No / Order ID / Location / Refund Amount / GST ... and a second block to the
    right; Oct 1-13 has just Order ID / Location / Amount / Payment gateway / Date / Ref / Reason."""
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        for hdr_row, r in enumerate(ws.iter_rows(min_row=1, max_row=3, values_only=True), start=1):
            header = [str(c or '').strip().lower() for c in r]
            if 'order id' in header and 'reason for refund' in header:
                break
        else:
            continue
        break
    else:
        raise ValueError(f'{path}: no sheet with an Order ID / Reason for Refund header')
    i_order, i_loc = header.index('order id'), header.index('location')
    i_amt = header.index('refund amount') if 'refund amount' in header else header.index('amount')
    i_reason, i_date, i_gw = header.index('reason for refund'), header.index('date'), header.index('payment gateway')
    rows = []
    for r in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
        if not r[i_order] or not r[i_amt]:
            continue
        rows.append({'order': str(r[i_order]).strip(), 'loc': str(r[i_loc] or '').strip(), 'amt': float(r[i_amt]),
                     'date': r[i_date], 'reason': norm_reason(r[i_reason]), 'gateway': r[i_gw]})
    return rows, wb


def cx_categories(path):
    """order -> {category: amount} from the CX refund tracker's refund sheet."""
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet, _ = find_refund_and_cn_sheets(wb)
    out = defaultdict(lambda: defaultdict(float))
    for r in wb[sheet].iter_rows(min_row=2, values_only=True):
        if r[2] and isinstance(r[7], (int, float)) and r[7]:
            out[str(r[2]).strip()][(r[6] or 'Uncategorized').strip()] += r[7]
    return out


def sales_categories(path, orders):
    """order -> ({category: weight}, POS location) from the sales history, for the given orders."""
    df = pd.read_csv(path, usecols=['Order name', 'Order or return', 'POS location name', 'Product type',
                                    'Gross sales', 'Taxes'])
    df = df[df['Order name'].isin(orders) & (df['Order or return'] == 'order')]
    out = {}
    for oid, g in df.groupby('Order name'):
        online = g['POS location name'].isna().all()
        w = defaultdict(float)
        for _, x in g.iterrows():
            pair = PRODUCT_TYPE_TO_REFUND_CAT.get(x['Product type'])
            cat = (pair[1] if online else pair[0]) if pair else 'Other'
            w[cat] += max(float(x['Gross sales'] or 0) + float(x['Taxes'] or 0), 0)
        if sum(w.values()) == 0:
            w = {k: 1.0 for k in w}
        out[oid] = (dict(w), g['POS location name'].dropna().iloc[0] if not online else None)
    return out


def build_month(rows, cx, sales):
    m = {'total': 0, 'onlineTotal': 0, 'storeTotal': 0, 'lineItems': 0, 'orders': set(),
         'byReason': defaultdict(float), 'byCategory': defaultdict(float), 'byStore': defaultdict(float),
         'byStoreCount': defaultdict(int), 'byReasonOnline': defaultdict(float),
         'byReasonOffline': defaultdict(float), 'byCategoryOnline': defaultdict(float),
         'byCategoryOffline': defaultdict(float)}
    src = Counter()
    for r in rows:
        store = LOCATION_MAP.get(r['loc'])
        if store is None and r['loc'] and not REGION_CODE.match(r['loc']):
            store = r['loc']  # newer ledgers name some stores plainly, e.g. "Oberoi Sky City"
        if store is None:
            store = (sales.get(r['order']) or ({}, None))[1]
            if store is None:
                raise ValueError(f"unknown ledger location {r['loc']!r} (order {r['order']}); add it to LOCATION_MAP")
        online = store == 'Online Store'
        if r['order'] in cx:
            weights, src['cx'] = cx[r['order']], src['cx'] + 1
        elif r['order'] in sales:
            weights, src['sales'] = sales[r['order']][0], src['sales'] + 1
        else:
            weights, src['none'] = {'Uncategorized': 1.0}, src['none'] + 1
        wsum = sum(weights.values())
        amt = r['amt']
        m['total'] += amt
        m['onlineTotal' if online else 'storeTotal'] += amt
        m['lineItems'] += 1
        m['orders'].add(r['order'])
        m['byReason'][r['reason']] += amt
        m['byStore'][store] += amt
        m['byStoreCount'][store] += 1
        m['byReasonOnline' if online else 'byReasonOffline'][r['reason']] += amt
        for cat, w in weights.items():
            share = amt * w / wsum
            m['byCategory'][cat] += share
            m['byCategoryOnline' if online else 'byCategoryOffline'][cat] += share
    out = {}
    for k, v in m.items():
        if k == 'orders':
            out[k] = len(v)
        elif isinstance(v, dict):
            out[k] = {kk: round(vv, 2) for kk, vv in v.items()} if not k.endswith('Count') else dict(v)
        else:
            out[k] = round(v, 2)
    return out, src


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--html', required=True)
    ap.add_argument('--ledger', required=True)
    ap.add_argument('--cx', help='CX refund tracker (refund + CN sheets) to take line-item categories from')
    ap.add_argument('--sales-history', required=True, help='Sales CSV covering the refunded orders\' purchase dates')
    ap.add_argument('--month', required=True, help='e.g. Oct-2026')
    a = ap.parse_args()

    rows, ledger_wb = load_ledger(a.ledger)
    months = Counter(month_key_from_date(r['date'], a.month) or a.month for r in rows)
    if set(months) != {a.month}:
        sys.exit(f'STOP: ledger has rows outside {a.month}: {dict(months)}')
    cx = cx_categories(a.cx) if a.cx else {}
    sales = sales_categories(a.sales_history, {r['order'] for r in rows if r['order'] not in cx}
                             | {r['order'] for r in rows if r['loc'] not in LOCATION_MAP})
    month, src = build_month(rows, cx, sales)

    with open(a.html, encoding='utf-8') as f:
        content = f.read()
    seed = extract_const(content, 'SEED_REFUNDS')
    seed['byMonth'][a.month] = month
    content = replace_const(content, 'SEED_REFUNDS', json.dumps(seed))
    # From Oct 1-13 the ledger file also carries the CX credit-note sheet; load it when present.
    _, cn_sheet = find_refund_and_cn_sheets(ledger_wb)
    cn = None
    if cn_sheet:
        cn_by_month = process_refund_sheet(ledger_wb[cn_sheet], amount_idx=7, reason_idx=8, date_idx=9,
                                           current_month=a.month)
        cn = cn_by_month.get(a.month)
        if cn:
            cn_seed = extract_const(content, 'SEED_CN')
            cn_seed['byMonth'][a.month] = cn
            content = replace_const(content, 'SEED_CN', json.dumps(cn_seed))
    with open(a.html, 'w', encoding='utf-8') as f:
        f.write(content)
    syntax_check_html_js(a.html)  # raises if the page's JS no longer parses

    print(f"OK. SEED_REFUNDS[{a.month}]: Rs{month['total']:,.0f} over {month['orders']} orders "
          f"({month['lineItems']} payouts); online Rs{month['onlineTotal']:,.0f}, stores Rs{month['storeTotal']:,.0f}")
    if cn:
        print(f"  SEED_CN[{a.month}] from the ledger's CN sheet: Rs{cn['total']:,.0f} over {cn['orders']} orders")
    print(f"  categories from CX tracker: {src['cx']} payouts, from sales lines: {src['sales']}, none: {src['none']}")
    dup = Counter(r['order'] for r in rows)
    for oid, n in dup.items():
        if n > 1:
            print(f"  NOTE {oid}: {n} payouts " + ', '.join(f"Rs{r['amt']:,.0f} {r['gateway']}" for r in rows if r['order'] == oid))


if __name__ == '__main__':
    main()
