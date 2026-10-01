#!/usr/bin/env python3
"""
month_archive.py -- per-month data files behind index.html's month dropdown.

index.html carries only the live month inline (SEED_DAYS etc.). Every closed month lives in
data/months/<YYYY-MM>.json with the same pieces the page needs to render it:

  {month, days (SEED_DAYS shape), baseline (BASELINE shape: prevMonth* = the month before,
   aug25/aug25StoreRevenue = same month last year), dailyTargets, targetsNote,
   refunds, cn (SEED_REFUNDS/SEED_CN shape), prevRefundCn (PREV_MONTH_REFUND_CN shape), nps}

and index.html's ARCHIVE_MONTHS const lists them. The page fetches a file when its month is
picked, so index.html doesn't grow by ~1MB per month.

  snapshot  Archive the live month straight out of index.html. Run this at every month
            transition, BEFORE update_main_dashboard.py --new-month overwrites it (and after
            the closing month's final refund/CN file has been loaded):
              month_archive.py snapshot --html index.html

  build     Build a past month from source files (used for the Apr-Aug 2026 backfill):
              month_archive.py build --month Jul-2026 --sales <csv covering the month AND the
                month before> [--targets <daywise xlsx> | --monthly-targets <Q1-style xlsx>]
                [--refund <xlsx>] [--prev-refund <xlsx>] [--from-git <commit>] --html index.html
            --from-git takes DAILY_TARGETS / SEED_REFUNDS / SEED_CN / SEED_NPS /
            PREV_MONTH_REFUND_CN from index.html as it was at that commit (the month's last
            refresh), overriding the file flags.
"""
import argparse
import calendar
import datetime as dt
import json
import os
import subprocess
import sys

import openpyxl
import pandas as pd

from bsc_common import (
    MONTHS, build_seed_days, extract_const, find_refund_and_cn_sheets, last_year_baseline,
    prepare_sales_df, process_refund_sheet, replace_const, syntax_check_html_js,
)
from update_main_dashboard import build_prev_month_baselines

OUT_DIR = os.path.join('data', 'months')
NODE = r'C:\Program Files\nodejs\node.exe' if os.name == 'nt' else 'node'


def month_parts(mk):
    mon, year = mk.split('-')
    return int(year), MONTHS.index(mon) + 1


def prev_month_key(mk):
    y, m = month_parts(mk)
    y, m = (y - 1, 12) if m == 1 else (y, m - 1)
    return '%s-%d' % (MONTHS[m - 1], y)


def file_for(mk):
    y, m = month_parts(mk)
    return '%d-%02d.json' % (y, m)


def eval_js_const(content, name):
    """BASELINE is a JS literal (comments, unquoted keys), not JSON -- let node evaluate it."""
    start = content.index('const %s = ' % name)
    end = content.index('\n};', start) + 3
    js = content[start:end] + '\nprocess.stdout.write(JSON.stringify(%s));' % name
    out = subprocess.run([NODE, '-e', js], capture_output=True, text=True, encoding='utf-8')
    if out.returncode != 0:
        raise RuntimeError(out.stderr)
    return json.loads(out.stdout)


def only_month(seed, mk):
    """Keep just one month's entry of a SEED_REFUNDS/SEED_CN-style {byMonth: {...}} blob."""
    if not seed or mk not in seed.get('byMonth', {}):
        return None
    return {'byMonth': {mk: seed['byMonth'][mk]}}


# ---------------------------------------------------------------- loaders

def read_month_slices(path, months):
    """Rows of a (possibly multi-year) sales export for the given 'YYYY-MM' prefixes, read in
    chunks so a 250MB file doesn't need to fit in memory whole."""
    keep = []
    for chunk in pd.read_csv(path, chunksize=200_000, low_memory=False):
        chunk = chunk[chunk['Day'].astype(str).str[:7].isin(months)]
        if len(chunk):
            keep.append(chunk)
    return pd.concat(keep) if keep else None


def load_daily_targets(path, mk):
    """Day-wise target workbook -> {store: {date: target}}. Handles both layouts seen so far:
    'POS location name, Date, New Revenue, Repeat Revenue' (Aug/Sep-26) and
    'Date (dd-mm-yy text), Store Name, Revenue, ...' (Jun-26)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else '' for c in next(rows)]
    col = {h: i for i, h in reversed(list(enumerate(header)))}
    store_i = col.get('POS location name', col.get('Store Name'))
    date_i = col['Date']
    y, m = month_parts(mk)
    out = {}
    for r in rows:
        store, d = r[store_i], r[date_i]
        if store is None or d is None:
            continue
        if isinstance(d, str):
            d = dt.datetime.strptime(d.strip(), '%d-%m-%y')
        if (d.year, d.month) != (y, m):
            continue
        if 'Revenue' in col:
            tgt = r[col['Revenue']] or 0
        else:
            tgt = (r[col['New Revenue']] or 0) + (r[col['Repeat Revenue']] or 0)
        out.setdefault(str(store).strip(), {})[d.strftime('%Y-%m-%d')] = round(tgt, 2)
    return out


# Store spellings in the Q1-2026 target workbook that differ from the POS names. "Oberoi Mall - New"
# = Oberoi Sky City is inferred (the only Oberoi store trading in Q1 without another target row).
Q1_STORE_MAP = {
    'Ambience VK': 'Ambience Vasant Kunj', 'DLF Midtown': 'DLF Midtown - Moti Nagar',
    'Inorbit Hyderabad': 'Inorbit mall Hyderabad', 'MOI': 'Mall of India, Noida',
    'Pali Hill, Khar': 'Pali Hill, Bandra', 'Vegas Mall Dwarka': 'Vegas Dwarka',
    'Oberoi Mall - New': 'Oberoi Sky City',
}


def load_monthly_targets(path, mk):
    """Q1-style workbook (Region, POS, <month-end date columns>...) -> month target per store,
    spread evenly over the month's days. Only month totals exist, so per-day figures are flat.
    The sheet holds an original table and, to its right, a revised one -- the revised one wins."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    y, m = month_parts(mk)
    header = rows[0]
    pos_i = max(i for i, c in enumerate(header) if c == 'POS')
    month_i = next(i for i in range(pos_i + 1, pos_i + 5)
                   if isinstance(header[i], dt.datetime) and (header[i].year, header[i].month) == (y, m))
    n = calendar.monthrange(y, m)[1]
    out = {}
    for r in rows[1:]:
        store, tgt = r[pos_i], r[month_i]
        if not store or not isinstance(tgt, (int, float)):
            continue
        store = str(store).strip()
        store = Q1_STORE_MAP.get(store, store)
        per_day = round(tgt / n, 2)
        out[store] = {'%d-%02d-%02d' % (y, m, d): per_day for d in range(1, n + 1)}
    return out


class PaddedSheet:
    """Older refund exports lack the trailing 'Date Processed' column process_refund_sheet
    reads; pad rows so a missing date falls back to the month being built."""
    def __init__(self, ws, width=12):
        self.ws, self.width = ws, width

    def iter_rows(self, **kw):
        for r in self.ws.iter_rows(**kw):
            yield tuple(r) + (None,) * max(0, self.width - len(r))


def load_refund_cn(path, mk):
    wb = openpyxl.load_workbook(path, data_only=True)
    rs, cs = find_refund_and_cn_sheets(wb)
    if rs is None or cs is None:
        raise ValueError('%s: could not find refund/CN sheets (%s)' % (path, wb.sheetnames))
    r = process_refund_sheet(PaddedSheet(wb[rs]), amount_idx=7, reason_idx=9, date_idx=10, current_month=mk)
    c = process_refund_sheet(PaddedSheet(wb[cs]), amount_idx=7, reason_idx=8, date_idx=9, current_month=mk)
    return only_month({'byMonth': r}, mk), only_month({'byMonth': c}, mk)


def git_index(commit):
    return subprocess.run(['git', 'show', commit + ':index.html'], capture_output=True,
                          text=True, encoding='utf-8', check=True).stdout


# ---------------------------------------------------------------- write + register

def write_bundle(bundle, html):
    os.makedirs(OUT_DIR, exist_ok=True)
    path = os.path.join(OUT_DIR, file_for(bundle['month']))
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(bundle, f, separators=(',', ':'))
    days = bundle['days']
    total = sum(sum(d['stores'].values()) for d in days.values())
    print('wrote %s: %d days, Rs%s revenue, targets=%s, refunds=%s, cn=%s, nps=%s' % (
        path, len(days), format(round(total), ','), bool(bundle['dailyTargets']),
        bool(bundle['refunds']), bool(bundle['cn']), bool(bundle['nps'])))
    register(html)


def register(html):
    """Rewrite index.html's ARCHIVE_MONTHS from whatever is in data/months/."""
    entries = []
    for fn in sorted(os.listdir(OUT_DIR)):
        if fn.endswith('.json'):
            y, m = int(fn[:4]), int(fn[5:7])
            entries.append({'key': '%s-%d' % (MONTHS[m - 1], y), 'file': 'data/months/' + fn})
    with open(html, encoding='utf-8') as f:
        content = f.read()
    content = replace_const(content, 'ARCHIVE_MONTHS', json.dumps(entries))
    with open(html, 'w', encoding='utf-8') as f:
        f.write(content)
    syntax_check_html_js(html)
    print('ARCHIVE_MONTHS: ' + ', '.join(e['key'] for e in entries))


# ---------------------------------------------------------------- commands

def cmd_snapshot(a):
    with open(a.html, encoding='utf-8') as f:
        content = f.read()
    days = extract_const(content, 'SEED_DAYS')
    last = max(days)
    mk = '%s-%s' % (MONTHS[int(last[5:7]) - 1], last[:4])
    prev = extract_const(content, 'PREV_MONTH_REFUND_CN')
    bundle = {
        'month': mk,
        'days': days,
        'baseline': eval_js_const(content, 'BASELINE'),
        'dailyTargets': extract_const(content, 'DAILY_TARGETS'),
        'targetsNote': None,
        'refunds': only_month(extract_const(content, 'SEED_REFUNDS'), mk),
        'cn': only_month(extract_const(content, 'SEED_CN'), mk),
        'prevRefundCn': prev if prev.get('month') == prev_month_key(mk) else None,
        'nps': extract_const(content, 'SEED_NPS'),
    }
    write_bundle(bundle, a.html)


def cmd_build(a):
    mk, pk = a.month, prev_month_key(a.month)
    y, m = month_parts(mk)
    py, pm = month_parts(pk)
    cur_p, prev_p = '%d-%02d' % (y, m), '%d-%02d' % (py, pm)
    print('reading %s for %s and %s...' % (a.sales, cur_p, prev_p))
    raw = read_month_slices(a.sales, {cur_p, prev_p})
    df = prepare_sales_df(raw)
    cur = df[df['Day_str'].str[:7] == cur_p]
    prv = df[df['Day_str'].str[:7] == prev_p]
    if cur.empty:
        sys.exit('STOP: no %s rows in %s' % (cur_p, a.sales))
    n_days = calendar.monthrange(y, m)[1]
    if cur['Day_str'].nunique() < n_days:
        print('WARNING: %s has only %d of %d days in this file' % (mk, cur['Day_str'].nunique(), n_days))

    baseline = build_prev_month_baselines(prv) if not prv.empty else {}
    ly, ly_stores = last_year_baseline(mk)
    baseline['aug25'] = ly
    baseline['aug25StoreRevenue'] = ly_stores

    targets, note = None, None
    if a.targets:
        targets = load_daily_targets(a.targets, mk)
    elif a.monthly_targets:
        targets = load_monthly_targets(a.monthly_targets, mk)
        note = ('Only month-level targets exist for %s, so each store\u2019s month target is spread '
                'evenly across the days; month totals and ach%% are exact, single-day targets are flat.' % mk)
    refunds = cn = prev_rc = nps = None
    if a.refund:
        refunds, cn = load_refund_cn(a.refund, mk)
    if a.prev_refund:
        pr, pc = load_refund_cn(a.prev_refund, pk)
        if pr or pc:
            prev_rc = {'month': pk, 'refund': (pr or {'byMonth': {pk: {}}})['byMonth'][pk],
                       'cn': (pc or {'byMonth': {pk: {}}})['byMonth'][pk]}
    if a.from_git:
        g = git_index(a.from_git)
        targets = extract_const(g, 'DAILY_TARGETS')
        note = None
        refunds = only_month(extract_const(g, 'SEED_REFUNDS'), mk)
        cn = only_month(extract_const(g, 'SEED_CN'), mk)
        nps = extract_const(g, 'SEED_NPS')
        try:
            p = extract_const(g, 'PREV_MONTH_REFUND_CN')
            if p.get('month') == pk:
                prev_rc = p
        except ValueError:
            pass
    if not targets:
        note = 'No target file on record for %s \u2014 actuals only, achievement %% not available.' % mk

    bundle = {'month': mk, 'days': build_seed_days(cur), 'baseline': baseline,
              'dailyTargets': targets or {}, 'targetsNote': note,
              'refunds': refunds, 'cn': cn, 'prevRefundCn': prev_rc, 'nps': nps}
    write_bundle(bundle, a.html)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest='cmd', required=True)
    s = sub.add_parser('snapshot')
    s.add_argument('--html', default='index.html')
    b = sub.add_parser('build')
    b.add_argument('--month', required=True, help='e.g. Jul-2026')
    b.add_argument('--sales', required=True, help='sales export covering the month and the month before it')
    b.add_argument('--targets')
    b.add_argument('--monthly-targets')
    b.add_argument('--refund')
    b.add_argument('--prev-refund')
    b.add_argument('--from-git')
    b.add_argument('--html', default='index.html')
    a = ap.parse_args()
    cmd_snapshot(a) if a.cmd == 'snapshot' else cmd_build(a)


if __name__ == '__main__':
    main()
